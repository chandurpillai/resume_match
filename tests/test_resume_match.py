import io
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from resume_match import COLLECTION, chunks, connect, details, export_excel, extract_text, ingest, rank_candidates


class TestEmbedder:
    """Deterministic vectors test plumbing without downloading model weights."""
    def embed(self, texts):
        for text in texts:
            vector = np.zeros(384)
            for i, word in enumerate(["python", "sales", "design"]):
                vector[i] = text.lower().count(word)
            vector[3] = 0.01
            yield vector / np.linalg.norm(vector)

    query_embed = embed


def test_persistent_ranking_and_duplicates(tmp_path):
    client = connect(path=str(tmp_path / "db"))
    embedder = TestEmbedder()
    resumes = [
        ("engineer.txt", "Alex Sample\nalex@example.com\nPython developer building Python services and APIs with testing and cloud deployment."),
        ("seller.txt", "Jordan Sample\njordan@example.com\nSales leader managing sales accounts and sales teams with excellent client relationship skills."),
        ("designer.txt", "Taylor Sample\nDesign specialist building design systems and design prototypes for accessible digital products and websites."),
    ]
    for filename, text in resumes:
        assert ingest(client, embedder, filename, text.encode())
    assert not ingest(client, embedder, "copy.txt", resumes[0][1].encode())
    assert client.count(COLLECTION).count == 3
    rows = rank_candidates(client, embedder, "Python engineer", 3, min_similarity=-1)
    assert [row["rank"] for row in rows] == [1, 2, 3]
    assert rows[0]["filename"] == "engineer.txt"
    assert rows[0]["email"] == "alex@example.com"
    assert "[Chunk 1]" in rows[0]["evidence"]
    assert rank_candidates(client, embedder, "Sales leader")[0]["filename"] == "seller.txt"
    client.close()
    reopened = connect(path=str(tmp_path / "db"))
    assert len(rank_candidates(reopened, embedder, "Design", 10, min_similarity=-1)) == 3
    reopened.close()


def test_empty_and_invalid_input(tmp_path):
    client = connect(path=str(tmp_path / "db"))
    assert rank_candidates(client, TestEmbedder(), "Python") == []
    with pytest.raises(ValueError):
        rank_candidates(client, TestEmbedder(), " ")
    client.close()
    for filename, data in [("empty.txt", b""), ("scan.txt", b"hello"), ("wrong.exe", b"sample")]:
        with pytest.raises(ValueError):
            extract_text(filename, data)


def test_custom_collection_and_incompatible_schema(tmp_path):
    from qdrant_client import QdrantClient, models
    path = str(tmp_path / "db")
    client = QdrantClient(path=path)
    client.create_collection("wrong_size", vectors_config=models.VectorParams(size=1536, distance=models.Distance.COSINE))
    client.close()
    with pytest.raises(ValueError, match="384-dimensional"):
        connect(path=path, collection="wrong_size")
    client = connect(path=path, collection="correct_size")
    assert client.collection_exists("wrong_size")
    embedder = TestEmbedder()
    text = b"Alex Sample Python developer building Python services and APIs with testing and cloud deployment."
    assert ingest(client, embedder, "sample.txt", text, collection="correct_size")
    assert rank_candidates(client, embedder, "Python", collection="correct_size")[0]["filename"] == "sample.txt"
    client.close()


def test_docx_tables_and_contact_extraction():
    doc = Document()
    doc.add_paragraph("Alex Sample")
    doc.add_paragraph("alex@example.com +1 312-555-0100 linkedin.com/in/alex-sample")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Python developer with experience in software testing and building production services"
    buffer = io.BytesIO()
    doc.save(buffer)
    text = extract_text("resume.docx", buffer.getvalue())
    assert "production services" in text
    result = details(text, "resume.docx")
    assert result["name"] == "Alex Sample"
    assert result["phone"] == "+1 312-555-0100"
    assert chunks(" ".join(str(i) for i in range(500)))[-1].endswith("499")


def test_filter_unrelated_candidates_by_similarity(tmp_path):
    client = connect(path=str(tmp_path / "filter_db"))
    model = TestEmbedder()
    ingest(client, model, "python.txt", b"Alex Sample Python developer building Python APIs with software testing and cloud deployment experience.")
    ingest(client, model, "sales.txt", b"Jordan Sample Sales manager leading sales teams and managing enterprise accounts and customer relationships.")
    assert rank_candidates(client, model, "civil engineer") == []
    rows = rank_candidates(client, model, "Python", limit=3)
    assert len(rows) == 1 and rows[0]["rank"] == 1
    assert len(rank_candidates(client, model, "Python", min_similarity=-1)) == 2
    client.close()


def test_excel_safe_strings_and_numeric_rank():
    data = export_excel([{"rank": 1, "name": "=HYPERLINK(\"evil\")", "score": 0.75,
                          "phone": "+1 312-555-0100", "evidence": "[Chunk 1] Python"}], "Python role")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        assert not root.findall(".//s:f", ns)
        assert root.find(".//s:c[@r='A2']/s:v", ns).text == "1"
        assert root.find(".//s:c[@r='G2']/s:v", ns).text == "0.75"
        assert root.find("s:autoFilter", ns) is not None
        assert "Job and methodology" in archive.read("xl/workbook.xml").decode()


def test_app_empty_job_validation():
    app = AppTest.from_file("../app.py", default_timeout=20).run()
    assert not app.exception
    next(button for button in app.button if button.label == "Rank candidates").click().run()
    assert not app.exception
    assert "Enter a job description" in app.error[0].value
