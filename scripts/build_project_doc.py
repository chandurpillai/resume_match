from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('outputs/Resume_match_project_guide.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(28)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(13)
doc.core_properties.title = 'Resume match project guide'
doc.core_properties.subject = 'Setup and operation of the Streamlit resume matching application'
doc.core_properties.author = ''

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def page(title):
    doc.add_page_break()
    doc.add_heading(title,1)
def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(9)
    for i,line in enumerate(text.splitlines()):
        run=para.add_run(('\n' if i else '')+line)
        run.font.name='Consolas'; run.font.size=Pt(9)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width= Inches(w)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        for cell,value,w in zip(row.cells,values,widths):
            cell.width=Inches(w); cell.text=value
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DDE7EF' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF')); pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'95'); el.set(qn('w:type'),'dxa'); margins.append(el)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.font.size=Pt(10); run.bold=i==0
        if i==0:
            header=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(header)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_paragraph('Resume match project guide','Title')
doc.add_paragraph('Streamlit resume retrieval and candidate review application','Subtitle')
p('This guide explains how to install, configure, use, and maintain the Resume match project in C:\\Dev\\Resume_match. It is intended for the developer running the application and the reviewer evaluating candidate resumes.')
p('The application stores resume embeddings in Qdrant, retrieves candidates against a job description, filters weak matches, and exports candidate details to Excel. Optional OpenAI reviews explain the retrieved evidence. Candidate rankings support human review and do not make hiring decisions.')
h('Capabilities')
table(['Feature','Implemented behavior'],[
['Resume intake','Multiple PDF, DOCX, and UTF-8 TXT uploads; up to 30 files per batch and 5 MB per file.'],
['Vector storage','Persistent embedded Qdrant or a configured Docker, server, or Cloud instance.'],
['Candidate matching','Maximum of three results by default; minimum similarity and optional required phrase filters.'],
['Evidence review','Retrieved excerpts plus optional OpenAI-generated qualifications, gaps, and interview questions.'],
['Excel download','Rank, contact details, source filename, score, evidence, candidate ID, and matching settings.']], [1.55,5.35])
h('Technology roles')
p('Streamlit provides the interface. FastEmbed runs the BAAI/bge-small-en-v1.5 embedding model locally. Qdrant stores and searches 384-dimensional vectors using Cosine similarity. The OpenAI Responses API generates reviews; it does not create this app’s embeddings. XlsxWriter builds the Excel download.')
h('Guide contents')
p('The following pages cover architecture and matching, Windows and Docker setup, daily operation and exports, and troubleshooting and validation.')

page('Architecture and matching')
h('Ingestion workflow')
p('Upload files → extract text and contact details → split into overlapping chunks → generate local embeddings → upsert vectors and payloads into Qdrant.')
p('PDF extraction uses pypdf, while DOCX extraction includes paragraphs and tables. Password-protected PDFs, PDFs over 50 pages, files with fewer than 10 readable words, and extracted text above 100,000 characters are rejected. Scanned PDFs require OCR before upload.')
p('Each chunk contains up to 120 words, with a 100-word stride. A SHA-256 hash of the extracted text identifies each resume; deterministic UUIDs identify its chunks. Re-uploading the same extracted text skips a complete existing record. Changed resume versions are separate candidates, even if the person is the same.')
h('Retrieval and ranking')
p('The job description is also chunked. Query embeddings are averaged and normalized, then Qdrant returns exact similarity scores for all stored chunks. Chunks are grouped by candidate. The candidate score is the mean of the best three chunk scores, or all available scores when fewer than three exist.')
code('candidate score = sum of best k chunk scores / k\nk = min(3, number of resume chunks)')
p('Candidates below the selected minimum similarity are excluded. The default cutoff is 0.65. Remaining candidates are sorted by descending score; a stable candidate ID resolves ties. Consecutive ranks are assigned only after filtering, up to the requested maximum. Zero, one, or two matches are valid results when the maximum is three.')
h('Phrase filtering and relevance')
p('The optional Required resume phrase field checks the stored resume chunks for a whole phrase, ignoring case and variable whitespace. Entering civil engineer requires that phrase; it does not accept a synonym such as structural designer. Semantic matching alone can return related experience without identical wording.')
p('A 0.65 cosine score is not a 65 percent qualification rating. The cutoff is a starting heuristic and should be tuned against reviewed examples. Resume length, broad job descriptions, and tokenizer truncation can affect relevance. Matching searches all chunks, so the implementation is intended for small candidate pools.')
h('Grounded review generation')
p('On request, OpenAI receives the submitted job description and the selected candidate’s best retrieved excerpts. The prompt requests chunk citations, supported qualifications, missing evidence, and interview questions. Requests use store=False and a 2,000-token output limit. Reviews can omit information outside these excerpts and must be checked against the resume.')

page('Windows and Docker setup')
h('Start the project')
p('Run PowerShell from the project folder. For a fresh environment, install Python 3.12 through uv and install the dependencies:')
code('cd C:\\Dev\\Resume_match\nuv venv .venv --python 3.12\nuv pip install --python .venv/Scripts/python.exe -r requirements.txt --link-mode copy')
p('The workspace already has a .venv. Its interpreter can launch Streamlit directly:')
code('.\\.venv\\Scripts\\python.exe -m streamlit run app.py')
p('Open http://localhost:8501. Initial embedding model download requires internet access. Model files are cached in data/models.')
h('Run Qdrant in Docker')
p('Start Docker Desktop with Linux containers. For a new container, use a named volume so stored data persists:')
code('docker volume create resume_qdrant_data\ndocker run -d --name resume-qdrant --restart unless-stopped `\n  -p 127.0.0.1:6333:6333 `\n  -v resume_qdrant_data:/qdrant/storage qdrant/qdrant')
p('If that container already exists, use docker start resume-qdrant. View collections at http://localhost:6333/dashboard. No API key is required by the shown local Docker configuration.')
h('Configure the application')
p('Create or edit .env in the project root. The app does not load .env.example as its runtime settings. Use this configuration, supplying your key locally:')
code('QDRANT_URL=http://localhost:6333\nQDRANT_API_KEY=\nQDRANT_PATH=data/qdrant\nQDRANT_COLLECTION=resume_match_bge384_app_v2\nOPENAI_API_KEY=your-api-key-here\nOPENAI_MODEL=gpt-4.1-mini')
p('Stop Streamlit with Ctrl+C and restart after configuration changes. QDRANT_PATH is used only when QDRANT_URL is empty. Keep actual credentials in .env, which is excluded from git.')
p('Let the app create its collection. The required schema is an unnamed vector of size 384 with Cosine distance. A collection name containing 384 does not enforce the vector size. An existing 1,536-dimensional collection is incompatible.')

page('Daily operation and Excel output')
h('Store and match resumes')
p('1. Upload supported resumes and select Store resumes. Review per-file errors and the stored and duplicate counts. Selecting a file alone does not store it.')
p('2. Enter a job description with responsibilities, skills, and experience. Set Maximum candidates to return and Minimum similarity. Optionally require a phrase such as civil engineer.')
p('3. Select Rank candidates. The app searches the selected Qdrant collection, including previously stored resumes. A result count below the maximum means fewer candidates passed the filters. An empty collection and a search with no matches produce different messages.')
p('4. Inspect candidate excerpts and verify extracted contact details. Use Generate evidence-based review when needed; this sends the job description and excerpts to OpenAI and incurs API usage charges.')
p('5. Select Download candidate details to save candidate_ranking.xlsx. Results and the download correspond to the last submitted search; submit again after changing the job or filters.')
h('Workbook contents')
table(['Worksheet','Contents'],[
['Candidate ranking','Rank, suggested name, email, phone, LinkedIn, resume filename, numeric cosine similarity, retrieved evidence, and candidate ID.'],
['Job and methodology','Submitted job description, score methodology, minimum similarity, required phrase, and interpretation notes.']], [1.7,5.2])
p('The ranking sheet includes a filter row, frozen headings, wrapped evidence, and numeric score formatting. Text is written as literal strings to prevent formula injection. Large evidence cells may need expansion in Excel.')
h('Data interpretation')
p('Names and contact fields are heuristic extractions rather than verified records. A blank field means it was not extracted. Phone parsing primarily supports common North American formats. Original upload files are not saved by the app, so retain the source resumes separately.')
h('Persistence and access')
p('Embedded Qdrant persists under data/qdrant. The Docker setup persists in the resume_qdrant_data volume. Streamlit restarts do not clear either database. Switching collections or connection settings selects another dataset; it does not migrate or delete the old dataset.')
p('The app has one shared collection and no built-in authentication or automatic resume deletion. Keep this local tool on localhost. Multi-user hosting requires access control, data isolation, and an appropriate retention process.')

page('Troubleshooting and validation')
h('Common fixes')
p('Python command not recognized: change to C:\\Dev\\Resume_match before using the relative .venv path. Verify that the environment exists and dependencies are installed.')
p('App still uses local storage: edit .env rather than .env.example, set QDRANT_URL=http://localhost:6333, and restart Streamlit. The sidebar should show Qdrant server and the selected collection.')
p('Incompatible collection: inspect the actual vector configuration in Qdrant. It must be unnamed, size 384, and Cosine. Select a new collection name in .env and restart so the app creates the correct schema. Existing resumes must be uploaded into the newly selected collection.')
p('Initialization fails: check that Docker is running and the Qdrant dashboard opens. Check model download connectivity and file access. The generic message can describe either storage or embedding initialization; schema errors now show a specific message.')
p('Unrelated candidates appear: increase the similarity cutoff, provide a more specific description, or add a required phrase. The phrase filter can exclude relevant synonyms. If new controls are missing, stop and restart the app from this project folder.')
p('OpenAI review fails: verify the API key in .env, available quota, model access, and network connectivity. Restart after changing settings. Ranking and Excel export do not require an OpenAI review.')
h('Validation commands and recorded checks')
code('.\\.venv\\Scripts\\python.exe -m pytest -q\n.\\.venv\\Scripts\\python.exe scripts/smoke_test.py')
p('The latest recorded automated run passed 11 tests. Coverage includes persistence, duplicate handling, custom collection validation, ranking and phrase filters, document parsing, Excel cell safety, Streamlit validation, and mocked OpenAI responses. A real embedding smoke test checked synthetic engineering and sales matches and rejected an unrelated civil engineer query. Live OpenAI generation was not verified.')
h('Project files')
p('app.py owns the Streamlit interface and resource caching. resume_match.py implements extraction, chunking, Qdrant operations, matching, OpenAI reviews, and Excel export. requirements.txt declares dependencies. tests contains automated checks. scripts/smoke_test.py exercises real embeddings. .streamlit/config.toml sets the local server and theme.')
h('Reference material')
p('Project implementation: app.py, resume_match.py, README.md, requirements.txt, and tests in C:\\Dev\\Resume_match. Configuration examples omit real credentials.')
p('Qdrant quickstart: https://qdrant.tech/documentation/quickstart/\nFastEmbed: https://qdrant.github.io/fastembed/\nStreamlit: https://docs.streamlit.io/\nOpenAI Responses API: https://developers.openai.com/api/reference/python/resources/responses/methods/create')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
